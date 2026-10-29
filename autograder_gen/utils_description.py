from io import BytesIO

from docx import Document

from autograder_gen.config import Config, DEFAULT_LANGUAGE_RUNTIMES
from autograder_gen.utils_submission import STUDENT_ID_FALLBACK, eval_expected_output_script


def generate_description_docx_buffer(config: Config) -> BytesIO:
    """Generate a Word document containing the assessment description."""
    doc = Document()
    doc.add_heading("Assessment Description", 0)
    runtime_info = DEFAULT_LANGUAGE_RUNTIMES.get(config.language.lower(), {})
    env_ver = runtime_info.get("version", config.language)
    doc.add_paragraph(f"Environment: {env_ver}")
    for i, question in enumerate(config.questions, 1):
        doc.add_heading(f"Question {i}: {question.name}", level=1)
        if hasattr(question, "description") and question.description:
            doc.add_paragraph(question.description)
        question_points = sum(item.total_mark for item in question.marking_items)
        p = doc.add_paragraph()
        run = p.add_run(f"Total Points: {question_points}")
        run.bold = True
        doc.add_heading("Marking Items", level=2)
        for visible_item_idx, item in enumerate(question.marking_items, 1):
            item_name = getattr(item, "name", "") or f"Marking Item {visible_item_idx}"
            doc.add_heading(f"{visible_item_idx}. {item_name}", level=3)
            doc.add_paragraph(f"Points: {item.total_mark}")
            if item.type == "program_output_test":
                doc.add_paragraph(
                    f"Requirement: Program must produce specific output for given input in '{item.target_file}' (output and expected output are trimmed before comparison)."
                )
                if item.expected_input:
                    doc.add_heading("Example Input:", level=4)
                    p = doc.add_paragraph()
                    run = p.add_run(item.expected_input)
                    run.font.name = "Courier New"
                exp_out = item.expected_output
                exp_script = getattr(item, "expected_output_script", None) or getattr(
                    question, "expected_output_script", None
                )
                if not exp_out and exp_script:
                    exp_out = eval_expected_output_script(
                        exp_script,
                        student_id=str(STUDENT_ID_FALLBACK),
                        expected_input=item.expected_input,
                    )
                if exp_out:
                    doc.add_heading("Expected Output:", level=4)
                    p = doc.add_paragraph()
                    run = p.add_run(exp_out)
                    run.font.name = "Courier New"
            elif item.type == "function_signature_test":
                sig_details = []
                if getattr(item, "expected_parameters", ""):
                    sig_details.append(f"parameters ({item.expected_parameters})")
                if getattr(item, "expected_return_type", ""):
                    sig_details.append(f"return type {item.expected_return_type}")
                detail_suffix = f" with {', '.join(sig_details)}" if sig_details else ""
                doc.add_paragraph(
                    f"Requirement: Function '{item.function_name}' in '{item.target_file}' must have correct signature{detail_suffix}."
                )
            elif item.type == "function_return_test":
                doc.add_paragraph(
                    f"Requirement: Function '{item.function_name}' in '{item.target_file}' must return expected values for given arguments."
                )
            elif item.type == "function_output_test":
                doc.add_paragraph(
                    f"Requirement: Function '{item.function_name}' in '{item.target_file}' must print expected standard output for given arguments."
                )
            elif item.type in ("git_submission_test", "git_submission_exists"):
                doc.add_paragraph(
                    "Requirement: Submission must be made via a Git repository (GitLab or GitHub)."
                )
            elif item.type in ("gitlab_submission_test", "gitlab_submission_exists"):
                doc.add_paragraph("Requirement: Submission must be made via a GitLab repository.")
            elif item.type in ("github_submission_test", "github_submission_exists"):
                doc.add_paragraph("Requirement: Submission must be made via a GitHub repository.")
            elif item.type == "manual_review":
                doc.add_paragraph(
                    f"Requirement: Qualitative evaluation of '{item.target_file}' (manual review by instructor)."
                    if item.target_file
                    else "Requirement: Qualitative evaluation (manual review by instructor)."
                )
            visible_item_idx += 1
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
